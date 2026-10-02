"""创建测试 docx 并转为 PDF，用于测试 PDF 上传场景。"""
from docx import Document

doc = Document()
doc.add_paragraph("张三 | 电话：13800000000 | 邮箱：zhangsan@example.com")

h = doc.add_paragraph("教育背景")
h.runs[0].bold = True
doc.add_paragraph("2013.09 - 2017.06 某大学 计算机科学与技术 本科")
doc.add_paragraph("主修课程：数据结构、操作系统、数据库原理")

h = doc.add_paragraph("实习经历")
h.runs[0].bold = True
doc.add_paragraph("2016.07 - 2016.09 某互联网公司 后端开发实习生")
doc.add_paragraph("负责后端接口的开发与测试。")
doc.add_paragraph("协助团队完成数据库迁移工作。")

h = doc.add_paragraph("校园经历")
h.runs[0].bold = True
doc.add_paragraph("2015.09 - 2016.06 计算机协会 技术部部长")
doc.add_paragraph("组织协会技术沙龙活动，参与人数约50人。")
doc.add_paragraph("协调部门成员完成活动策划与执行。")

h = doc.add_paragraph("专业技能")
h.runs[0].bold = True
doc.add_paragraph("熟悉Java、Python编程语言")
doc.add_paragraph("了解MySQL数据库和Redis缓存")

h = doc.add_paragraph("证书")
h.runs[0].bold = True
doc.add_paragraph("CET-6 大学英语六级")

h = doc.add_paragraph("自我评价")
h.runs[0].bold = True
doc.add_paragraph("本人性格开朗，善于沟通，有较强的学习能力和团队合作精神。")

doc.save("test_resume.docx")
print("docx created")

# 转为 PDF
from docx2pdf import convert
convert("test_resume.docx", "test_resume.pdf")
print("pdf created")
