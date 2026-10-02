"""ATS 文本层校验（免费功能，Phase 2）。

从 ai-job-search 的 verify_pdf.py 提炼的 ATS 文本层校验逻辑，移植为纯 Python：
- PDF：用 PyMuPDF(fitz) 抽取文本层，逐页统计可读文本量，识别「图片型空页」（ATS 读不到）。
- Word：用 python-docx 抽取段落 + 表格文本。
- 校验项：乱码字形(cid/替换符)、联系方式是否为字面文本、日期是否用了 ATS 易丢的 en-dash、
         图片型空页、JD 关键词覆盖率。

文件由前端导出后直接上传字节流，与 main.py 的 sessions 解耦，零外部状态依赖。
"""
import io
import re
from typing import Any, Dict, List, Optional

import fitz  # PyMuPDF
from docx import Document

from .profile_bridge import extract_jd_keywords, keyword_overlap

# ── 正则 ──
_RE_EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_RE_PHONE = re.compile(r"(?<!\d)(?:1[3-9]\d{9}|0\d{2,3}-?\d{7,8})(?!\d)")
# 年份之间出现 en-dash(U+2013)/em-dash(U+2014)/波浪号被替换的情况
_RE_DASH_DATE = re.compile(r"\d{4}\s*[\u2013\u2014]\s*(?:\d{4}|至今|现在|present|Present)", re.IGNORECASE)
# 乱码：PDF 中未嵌入字体时常见的 (cid:XX) 与 Unicode 替换符
_RE_CID = re.compile(r"\(cid:\d+\)")
_RE_REPL = re.compile(r"\ufffd")
_RE_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

_EMAIL_WORDS = ("email", "e-mail", "邮箱", "电子邮件", "mail")
_PHONE_WORDS = ("phone", "tel", "mobile", "电话", "手机", "联系", "contact")


def extract_pdf_text(stream: bytes) -> Dict[str, Any]:
    """返回 {full_text, pages:[{page,text,text_len}]}。"""
    doc = fitz.open(stream=stream, filetype="pdf")
    pages: List[Dict[str, Any]] = []
    full: List[str] = []
    for i, page in enumerate(doc):
        t = page.get_text("text") or ""
        pages.append({"page": i + 1, "text": t, "text_len": len(t.strip())})
        full.append(t)
    doc.close()
    return {"full_text": "\n".join(full), "pages": pages}


def extract_docx_text(stream: bytes) -> Dict[str, Any]:
    doc = Document(io.BytesIO(stream))
    texts: List[str] = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text and cell.text.strip():
                    texts.append(cell.text)
    full = "\n".join(texts)
    # Word 没有「页」概念，按字数近似分页供展示
    approx_pages = max(1, (len(full) + 1500) // 1500)
    pages = [{"page": i + 1, "text": "", "text_len": 0} for i in range(approx_pages)]
    return {"full_text": full, "pages": pages}


def _check_garbled(text: str) -> Dict[str, Any]:
    cids = _RE_CID.findall(text)
    repls = len(_RE_REPL.findall(text))
    ctrl = len(_RE_CTRL.findall(text))
    bad = len(cids) + repls + ctrl
    status = "pass" if bad == 0 else ("warn" if bad <= 3 else "fail")
    examples = []
    for m in _RE_CID.findall(text)[:3]:
        examples.append(m)
    if repls:
        examples.append("\ufffd")
    return {
        "name": "文本层无乱码",
        "status": status,
        "detail": (
            "字体已正常嵌入，文字可被 ATS 解析"
            if bad == 0
            else f"检测到 {bad} 处疑似乱码（cid/替换符/控制符），ATS 可能读成空白或方块"
        ),
        "examples": examples,
    }


def _check_image_only_pages(pages: List[Dict[str, Any]], fmt: str) -> Dict[str, Any]:
    if fmt != "pdf":
        return {"name": "页面非图片型", "status": "pass",
                "detail": "Word 为原生文本，无图片型空页风险", "pages_empty": 0}
    empty = [p["page"] for p in pages if p["text_len"] == 0]
    status = "pass" if not empty else ("warn" if len(empty) == 1 else "fail")
    return {
        "name": "无图片型空页",
        "status": status,
        "detail": (
            "每页均有可提取文本，ATS 能完整读取"
            if not empty
            else f"第 {empty} 页无可提取文本（疑似纯图片/扫描），ATS 将读不到该页内容"
        ),
        "pages_empty": empty,
    }


def _check_contact_literal(text: str, profile_text: Optional[str]) -> Dict[str, Any]:
    emails = _RE_EMAIL.findall(text)
    phones = _RE_PHONE.findall(text)
    # 若提供了档案，用档案里的联系方式核对是否出现在导出文本中
    if profile_text:
        p_emails = _RE_EMAIL.findall(profile_text)
        p_phones = _RE_PHONE.findall(profile_text)
        miss_e = [e for e in p_emails if e not in text]
        miss_p = [p for p in p_phones if p not in text]
        if p_emails and miss_e:
            return {"name": "联系方式为字面文本", "status": "fail",
                    "detail": f"档案中的邮箱 {miss_e} 未以字面文本出现在导出文件中，ATS 读不到",
                    "emails": emails, "phones": phones}
        if p_phones and miss_p:
            return {"name": "联系方式为字面文本", "status": "fail",
                    "detail": f"档案中的电话 {miss_p} 未以字面文本出现在导出文件中，ATS 读不到",
                    "emails": emails, "phones": phones}
        return {"name": "联系方式为字面文本", "status": "pass",
                "detail": "邮箱/电话均已以字面文本呈现，ATS 可抓取",
                "emails": emails, "phones": phones}
    # 未提供档案：仅报告是否检测到
    if emails or phones:
        return {"name": "联系方式为字面文本", "status": "pass",
                "detail": "检测到字面邮箱/电话，ATS 可抓取（未提供档案做强校验）",
                "emails": emails, "phones": phones}
    return {"name": "联系方式为字面文本", "status": "warn",
            "detail": "未在导出文本中识别到邮箱/电话字面，请确认联系方式已正常填写且非图标",
            "emails": emails, "phones": phones}


def _check_dash_dates(text: str) -> Dict[str, Any]:
    hits = _RE_DASH_DATE.findall(text)
    status = "pass" if not hits else "warn"
    return {
        "name": "日期未用易丢连接符",
        "status": status,
        "detail": (
            "日期区间使用标准格式，ATS 可正确解析"
            if not hits
            else f"发现 {len(hits)} 处 en-dash/em-dash 日期（如 2022–2023），部分 ATS 会静默截断，建议改用「~」或「至」"
        ),
        "examples": list(hits[:5]),
    }


def ats_check(file_bytes: bytes, fmt: str, jd: str = "", target_job: str = "",
              profile: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """对导出的 PDF/Word 文本层做 ATS 校验，返回统一报告。"""
    fmt = (fmt or "").lower()
    extracted = extract_pdf_text(file_bytes) if fmt == "pdf" else extract_docx_text(file_bytes)
    full_text = extracted["full_text"]
    pages = extracted["pages"]

    profile_text = None
    if isinstance(profile, dict):
        from .profile_bridge import flatten_profile_text
        profile_text = flatten_profile_text(profile)
        # flatten_profile_text 不含 contact 字段，显式补上以便强校验联系方式字面性
        basic = profile.get("basic") or {}
        for k in ("email", "phone", "mobile"):
            v = basic.get(k)
            if v:
                profile_text = (profile_text or "") + "\n" + str(v)

    checks: List[Dict[str, Any]] = []
    checks.append(_check_garbled(full_text))
    checks.append(_check_image_only_pages(pages, fmt))
    checks.append(_check_contact_literal(full_text, profile_text))
    checks.append(_check_dash_dates(full_text))

    # 结构性合规分：乱码/空页/联系方式是硬伤，连接符是软伤
    score = 100
    for c in checks:
        if c["name"] == "文本层无乱码" and c["status"] == "fail":
            score -= 35
        elif c["name"] == "文本层无乱码" and c["status"] == "warn":
            score -= 12
        elif c["name"] == "无图片型空页" and c["status"] == "fail":
            score -= 30
        elif c["name"] == "无图片型空页" and c["status"] == "warn":
            score -= 12
        elif c["name"] == "联系方式为字面文本" and c["status"] == "fail":
            score -= 18
        elif c["name"] == "联系方式为字面文本" and c["status"] == "warn":
            score -= 6
        elif c["name"] == "日期未用易丢连接符" and c["status"] == "warn":
            score -= 8
    score = max(0, min(100, score))

    # JD 关键词覆盖率（若有 JD）
    jd_coverage = None
    if jd or target_job:
        kws = extract_jd_keywords(jd, target_job)
        ov = keyword_overlap(kws, full_text)
        jd_coverage = {
            "coverage": ov["coverage"],
            "jd_keyword_count": ov["jd_keyword_count"],
            "matched": ov["matched"][:40],
            "missing": ov["missing"][:40],
        }
        score = round(score * 0.6 + ov["coverage"] * 0.4)

    if score >= 80:
        pass_level = "pass"
    elif score >= 60:
        pass_level = "warn"
    else:
        pass_level = "fail"

    verdict = {
        "pass": "ATS 兼容性良好，机器筛选可正常解析",
        "warn": "存在若干 ATS 兼容隐患，建议优化后再投递",
        "fail": "ATS 兼容性较差，可能被机器筛选直接淘汰",
    }[pass_level]

    return {
        "ats_score": score,
        "pass_level": pass_level,
        "verdict_text": verdict,
        "format": fmt,
        "page_count": len(pages) if fmt == "pdf" else None,
        "text_length": len(full_text),
        "checks": checks,
        "jd_coverage": jd_coverage,
        "method": "fitz+python-docx",
    }
