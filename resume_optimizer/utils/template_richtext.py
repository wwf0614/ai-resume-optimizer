# -*- coding: utf-8 -*-
"""新编辑模式转换引擎：模板骨架 + 就地富文本编辑。

把 Word 模板转换为「无字装饰底图 + 可编辑文字层」：
- build_richtext_template   一次构建：几何校准 + 字体格式 + 板块语义 +
                            照片框识别 + 无字装饰底图
- apply_rich_texts          编辑内容（含格式）写回模板 docx（导出用）
- replace_photo_bytes       按原图纵横比居中裁剪后替换模板照片

定位/几何复用 template_editor 的解析（EMU→px、分组坐标系、表格单元格、
PDF 实际渲染校准），写回时保留模板原格式；长文本自动切换 Word
「缩小字体填充」防跑版。
"""
import io
import re
import shutil
import tempfile
from pathlib import Path
from typing import Dict, List, Optional

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from . import template_editor as te
from .template_editor import (
    A_NS, MC_NS, WP_NS,
    _box_geometry, _collect_edit_targets, _emu, _find_ancestor,
    _node_text, _paragraph_info,
)

PIC_NS = "http://schemas.openxmlformats.org/drawingml/2006/picture"
EMU_PER_PX = 9525

# 基础信息字段标签 → 标准键（用于标签盒/值盒配对与内容跟随）
FIELD_LABEL_MAP = {
    "姓名": "name", "名字": "name",
    "性别": "gender", "年龄": "age", "生日": "birth_date", "出生日期": "birth_date",
    "出生年月": "birth_date", "籍贯": "hometown", "民族": "nation",
    "身高": "height", "体重": "weight", "婚姻状况": "marriage", "政治面貌": "political_status",
    "现居": "city", "现居地址": "city", "现居城市": "city", "所在城市": "city",
    "地址": "city", "住址": "city",
    "学历": "education_degree", "专业": "major", "毕业院校": "school", "毕业学校": "school",
    "手机": "phone", "手机号码": "phone", "电话": "phone", "联系电话": "phone",
    "邮箱": "email", "电子邮箱": "email", "微信": "wechat", "微博": "weibo", "QQ": "qq",
    "工作经验": "work_years", "工作年限": "work_years",
    "求职意向": "intention_job", "应聘岗位": "intention_job", "期望职位": "intention_job",
}


def _role_of_title(text: str) -> Optional[str]:
    """板块标题文本 → 标准角色键。"""
    t = (text or "").strip()
    if not t or len(t) > 10:
        return None
    t = re.sub(r"[/／].*$", "", t).strip()  # 「教育背景 / Education」取中文段
    if not t or len(t) > 8:
        return None
    if re.match(r"^(个人信息|基本(信息|资料)|关于我)", t):
        return "basic"
    if re.match(r"^(教育)", t):
        return "education"
    if re.match(r"^(工作|职业经历|职业履历|任职)", t):
        return "work"
    if re.match(r"^(实习)", t):
        return "internship"
    if re.match(r"^(项目|作品)", t):
        return "project"
    if re.match(r"^(校园|社会实践|校内|社团|学生会|干部)", t):
        return "campus"
    if re.match(r"^(技能|专业(技能|课程)|语言|计算机|掌握技能|IT)", t):
        return "skill"
    if re.match(r"^(荣誉|获奖|证书|奖项|奖状)", t):
        return "honor"
    if re.match(r"^(自我(评价|介绍)|个人优势|职业(优势|总结))", t):
        return "self"
    if re.match(r"^(兴趣|爱好|特长)", t):
        return "hobby"
    if re.match(r"^(其他|附加|更多)", t):
        return "custom"
    return None


def _para_fmt(p_elem) -> Dict:
    """段落对齐/行距（相对倍数）。"""
    p_pr = p_elem.find(qn("w:pPr"))
    align, line = "left", None
    if p_pr is not None:
        jc = p_pr.find(qn("w:jc"))
        if jc is not None and jc.get(qn("w:val")):
            align = jc.get(qn("w:val"))
        sp = p_pr.find(qn("w:spacing"))
        if sp is not None:
            v = sp.get(qn("w:line"))
            if v and (sp.get(qn("w:lineRule")) or "auto") == "auto":
                try:
                    line = round(int(v) / 240.0, 2)
                except ValueError:
                    line = None
    return {"align": align, "line": line}


def _box_text_props(txbx) -> Dict:
    """文本框 bodyPr：垂直锚点 + 内边距（px）。"""
    cur = txbx.getparent()
    while cur is not None:
        if cur.tag == "{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}wsp":
            bp = cur.find("{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}bodyPr")
            if bp is None:
                return {}
            def ins(name, default):
                v = bp.get(name)
                try:
                    return round(int(v) / EMU_PER_PX) if v else default
                except (TypeError, ValueError):
                    return default
            return {"anchor": bp.get("anchor") or "t",
                    "pl": ins("lIns", 8), "pr": ins("rIns", 8),
                    "pt": ins("tIns", 5), "pb": ins("bIns", 5)}
        cur = cur.getparent()
    return {}


def _cell_text_props(_tc) -> Dict:
    return {"anchor": "t", "pl": 8, "pr": 8, "pt": 5, "pb": 5}


def _first_font(paragraphs: List[Dict]) -> Dict:
    """取盒子内首个带文字 run 的字体作为盒子默认格式。"""
    for p in paragraphs:
        for r in p.get("runs") or []:
            if (r.get("text") or "").strip():
                return {"family": r.get("font") or "微软雅黑",
                        "size": r.get("size") or 12,
                        "color": r.get("color") or "#333333",
                        "bold": bool(r.get("bold")),
                        "italic": bool(r.get("italic"))}
    return {"family": "微软雅黑", "size": 12, "color": "#333333",
            "bold": False, "italic": False}


def _is_basic_field_box(text: str) -> bool:
    """盒子内含 ≥2 个「标签：」行 → 基础信息标签盒。"""
    hits = 0
    for label in FIELD_LABEL_MAP:
        if re.search(re.escape(label) + r"\s*[：:]", text):
            hits += 1
    return hits >= 2


def _single_field_box(text: str) -> bool:
    """单个字段盒：{{姓名}} 或「求职意向：xxx」这类一行式基础信息。"""
    t = (text or "").strip()
    if not t or len(t) > 30:
        return False
    phs = te.PLACEHOLDER_RE.findall(t)
    if len(phs) == 1 and len(phs[0]) <= 6 and phs[0] in FIELD_LABEL_MAP \
            and t.replace("{{" + phs[0] + "}}", "").strip() in ("", "：", ":"):
        return True
    for label in FIELD_LABEL_MAP:
        if t.startswith(label) and ("：" in t or ":" in t):
            return True
    return False


def _annotate_roles(boxes: List[Dict]) -> None:
    """给每个盒子标注 role/kind：
    kind: title（板块标题）/ content（板块内容）/ field（基础信息）/
          free（自由内容，无板块归属）/ deco（不进编辑层，留在底图）
    role: basic/education/work/... 或 None
    顺序：板块标题 → 基础信息（多标签盒/单字段盒/值盒伴生）→ 文档相邻
    归属（内容盒在 XML 中通常紧贴自己的标题盒）→ 几何兜底 → free。"""
    editable = [b for b in boxes if b.get("editable")]
    titles = []
    for b in editable:
        role = _role_of_title(b.get("text") or "")
        if role and len((b.get("text") or "").strip()) <= 10:
            b["kind"], b["role"] = "title", role
            titles.append(b)
    title_ids = {b["id"] for b in titles}

    # 基础信息：多标签盒 / 单字段盒 / 值盒伴生（只有多标签盒可种子伴生，
    # 避免把右侧的正常板块内容盒误吸进来）
    field_boxes = []
    seed_boxes = []
    for b in editable:
        if b.get("kind"):
            continue
        t = b.get("text") or ""
        if _is_basic_field_box(t):
            b["kind"], b["role"] = "field", "basic"
            field_boxes.append(b)
            seed_boxes.append(b)
        elif _single_field_box(t):
            b["kind"], b["role"] = "field", "basic"
            field_boxes.append(b)

    def _companion(b, fb) -> bool:
        same_band = b["y"] < fb["y"] + fb["h"] and b["y"] + b["h"] > fb["y"]
        if not same_band:
            return False
        right = fb["x"] + fb["w"] - 12 <= b["x"] <= fb["x"] + fb["w"] * 2 + 60
        if right:
            return True
        # 横向明显重叠（标签盒与值盒叠放的模板）
        overlap = min(b["x"] + b["w"], fb["x"] + fb["w"]) - max(b["x"], fb["x"])
        return overlap > 0.4 * min(b["w"], fb["w"]) and b["w"] >= fb["w"]

    grown = True
    while grown and field_boxes:
        grown = False
        for b in editable:
            if b.get("kind"):
                continue
            for fb in seed_boxes:
                if _companion(b, fb):
                    b["kind"], b["role"] = "field", "basic"
                    field_boxes.append(b)
                    grown = True
                    break

    owner_of: Dict[int, Dict] = {}
    for t in titles:
        idx = boxes.index(t)
        for j in (idx - 1, idx + 1):
            if 0 <= j < len(boxes):
                cand = boxes[j]
                if not cand.get("editable") or cand.get("kind") or cand["id"] in title_ids:
                    continue
                owner_of[cand["id"]] = t
                break
    for b in editable:
        if b.get("kind") or b["id"] in title_ids or b["id"] in owner_of:
            continue
        owner = None
        for t in titles:
            if t["y"] <= b["y"] + 6 and (owner is None or t["y"] > owner["y"]):
                owner = t
        if owner is not None:
            owner_of[b["id"]] = owner

    for b in editable:
        if b.get("kind"):
            continue
        owner = owner_of.get(b["id"])
        if owner is not None and b["w"] * b["h"] >= 120:
            b["kind"], b["role"] = "content", owner["role"]
        else:
            # 无板块归属的文字盒仍可编辑（自由内容），只是不进板块导航
            b["kind"], b["role"] = "free", None


def _detect_photos(body) -> List[Dict]:
    """识别锚定图片（照片候选）：位置/尺寸/形状/关系 id，按面积降序。"""
    out = []
    seen = set()
    for anchor in body.iter("{%s}anchor" % WP_NS):
        if _find_ancestor(anchor, "{%s}Fallback" % MC_NS) is not None:
            continue
        pics = anchor.findall(".//{%s}pic" % PIC_NS)
        if not pics:
            continue
        ext = anchor.find("{%s}extent" % WP_NS)
        w = _emu(ext.get("cx")) / EMU_PER_PX if ext is not None else 0
        h = _emu(ext.get("cy")) / EMU_PER_PX if ext is not None else 0
        if w < 50 or h < 50:
            continue
        pos_x = pos_y = 0
        pos_h = anchor.find("{%s}positionH" % WP_NS)
        pos_v = anchor.find("{%s}positionV" % WP_NS)
        if pos_h is not None:
            off = pos_h.find("{%s}posOffset" % WP_NS)
            if off is not None:
                pos_x = _emu(off.text) / EMU_PER_PX
        if pos_v is not None:
            off = pos_v.find("{%s}posOffset" % WP_NS)
            if off is not None:
                pos_y = _emu(off.text) / EMU_PER_PX
        blip = anchor.find(".//{%s}blip" % A_NS)
        rid = blip.get(qn("r:embed")) if blip is not None else None
        if not rid:
            continue
        prst_el = anchor.find(".//{%s}prstGeom" % A_NS)
        prst = prst_el.get("prst") if prst_el is not None else "rect"
        key = (rid, round(pos_x), round(pos_y))
        if key in seen:
            continue
        seen.add(key)
        out.append({"x": round(pos_x), "y": round(pos_y),
                    "w": round(w), "h": round(h),
                    "rid": rid, "shape": prst or "rect"})
    out.sort(key=lambda b: -(b["w"] * b["h"]))
    return out


def _to_pdf(docx_path, pdf_path, timeout: int = 60) -> bool:
    """快速 PDF 转换：优先复用全局 Word 进程（省 3-5 秒/次），
    失败自动退回独立子进程模式。"""
    try:
        from . import converter
        converter.docx_to_pdf_fast(str(docx_path), str(pdf_path))
        if Path(pdf_path).exists():
            return True
    except Exception:  # noqa: BLE001
        pass
    return te._convert_to_pdf_safe(docx_path, pdf_path, timeout=timeout)


def _collect_boxes(body) -> List[Dict]:
    """收集全部可编辑目标（文本框/单元格）的几何、文本与格式。
    纯隐藏文字（w:vanish）的盒子标记 hidden，后续不进编辑层。"""
    boxes = []
    for src_kind, node, cell_geo in _collect_edit_targets(body):
        paragraphs = []
        all_hidden = True
        has_text = False
        for p in node.findall(qn("w:p")):
            info = _paragraph_info(p)
            fmt = _para_fmt(p)
            if (info["text"] or "").strip():
                has_text = True
                # 该段是否存在任一非隐藏 run
                runs = []
                for r in info["runs"]:
                    runs.append(r)
                visible = _para_visible(p)
                if visible:
                    all_hidden = False
            else:
                continue
            paragraphs.append({"text": info["text"], "runs": info["runs"],
                               "align": fmt["align"], "line": fmt["line"]})
        full_text = "\n".join(p["text"] for p in paragraphs).strip("\n")
        if src_kind == "textbox":
            geo = _box_geometry(node)
            props = _box_text_props(node)
        else:
            geo = cell_geo
            props = _cell_text_props(node)
        boxes.append({
            "id": len(boxes), "src": src_kind,
            "x": geo["x"], "y": geo["y"], "w": geo["w"], "h": geo["h"],
            "text": full_text,
            "hidden": has_text and all_hidden,
            "paragraphs": paragraphs,
            "font": _first_font(paragraphs),
            "props": props,
            "placeholders": te.PLACEHOLDER_RE.findall(full_text),
        })
    return boxes


def _para_visible(p_elem) -> bool:
    """段落里是否至少有一个非隐藏(w:vanish)且带文字的 run。"""
    for r in p_elem.findall(qn("w:r")):
        t = "".join(x.text or "" for x in r.iter(qn("w:t")))
        if not t.strip():
            continue
        rpr = r.find(qn("w:rPr"))
        if rpr is not None and rpr.find(qn("w:vanish")) is not None:
            continue
        return True
    return False


def _pdf_hit(page, key):
    """递减长度的候选键搜索（长句首行截短，避免命中到行中其他列）。"""
    for n in (24, 16, 10, 6):
        cand = (key or "")[:n].strip()
        if len(cand) < 2:
            continue
        try:
            hits = page.search_for(cand)
            if hits:
                return hits
        except Exception:  # noqa: BLE001
            continue
    return []


def _calibrate_containment(pdf_path, boxes, ml: int, mt: int) -> set:
    """矩形包含式 PDF 校准：文本命中点落在哪个盒子的近似 XML 区域内，
    就把该盒子校准到命中位置（取区域内最顶部的命中=盒子首行）。
    比按顺序配对更可靠——同一模板常见「教育背景两条同名日期行」，
    顺序配对会整体错位一格。"""
    from collections import defaultdict

    try:
        import fitz
    except Exception:  # noqa: BLE001
        return set()
    scale = 96.0 / 72.0
    hit_ids = set()
    doc = fitz.open(str(pdf_path))
    try:
        page = doc[0]
        groups = defaultdict(list)
        for b in boxes:
            lines = [ln.strip() for ln in (b.get("text") or "").splitlines() if ln.strip()]
            if lines:
                groups[lines[0][:32]].append(b)
        for key, grp in groups.items():
            hits = _pdf_hit(page, key)
            for b in grp:
                bx, by = b["x"] + ml, b["y"] + mt
                best = None
                for r in hits:
                    cx = (r.x0 + r.x1) / 2 * scale
                    cy = (r.y0 + r.y1) / 2 * scale
                    if bx - 30 <= cx <= bx + b["w"] + 30 and by - 30 <= cy <= by + b["h"] + 30:
                        if best is None or cy < best[0]:
                            best = (cy, r)
                if best is None:
                    continue
                r = best[1]
                b["x"] = round(r.x0 * scale - 6)
                b["y"] = round(r.y0 * scale - 6)
                if not b["w"]:
                    b["w"] = round(r.width * scale)
                if not b["h"]:
                    b["h"] = round(max(r.height * scale, 20))
                hit_ids.add(b["id"])
    finally:
        try:
            doc.close()
        except Exception:  # noqa: BLE001
            pass
    return hit_ids


def _extract_para_anchors(pdf_path, boxes) -> None:
    """从 Word 渲染的 PDF 中提取每个盒子里每段落的实测 y 坐标（px，相对盒顶）
    与盒内实测行距。前端按锚点给每段做相对定位，消除 HTML 与 Word
    行高/换行差异导致的累积漂移（文字压到下一板块的问题）。"""
    try:
        import fitz
    except Exception:  # noqa: BLE001
        return
    scale = 96.0 / 72.0
    doc = fitz.open(str(pdf_path))
    try:
        page = doc[0]
        data = page.get_text("dict")
        rendered = []
        for blk in data.get("blocks", []):
            if blk.get("type") != 0:
                continue
            for ln in blk.get("lines", []):
                txt = "".join(sp.get("text", "") for sp in ln.get("spans", [])).strip()
                if not txt:
                    continue
                x0, y0, _x1, y1 = ln["bbox"]
                rendered.append({"text": txt, "x0": x0 * scale,
                                 "y0": y0 * scale, "y1": y1 * scale})
        rendered.sort(key=lambda r: (round(r["y0"], 1), r["x0"]))

        for b in boxes:
            if not b.get("editable"):
                continue
            rx0, ry0 = b["x"], b["y"]
            rx1, ry1 = b["x"] + b["w"], b["y"] + b["h"]
            # 只按纵向范围收集（盒子 x 取自首行缩进位置，正文行可能更靠左，
            # 横向差异由每段自己的 x 锚点修正）
            lines_in = [r for r in rendered
                        if ry0 - 4 <= (r["y0"] + r["y1"]) / 2 <= ry1 + 6]
            if not lines_in:
                continue

            # 段落锚点：按归一化文本前缀顺序匹配渲染行；行距按段落实测
            # （本段首行到下一段首行之间的渲染行间距），避免全盒中位数
            # 与个别段落的真实行距不符导致段尾文字压到下一板块。
            # Word 自动编号列表的渲染行带「1. 」前缀，匹配时需剥离。
            _num_re = re.compile(r"^(?:\d{1,3}[.、)．]?[\s　]*|[•·]\s*)")

            def _norm(t):
                t = re.sub(r"\s+", "", t)
                return _num_re.sub("", t)

            anchors = []
            pitches = []
            para_match = []
            ptr = 0
            total_lines = len(lines_in)
            for pi, p in enumerate(b.get("paragraphs") or []):
                ptxt = _norm(p.get("text", ""))
                if not ptxt:
                    anchors.append(None)
                    pitches.append(None)
                    continue
                matched = None
                idx = ptr
                while idx < total_lines:
                    lt = _norm(lines_in[idx]["text"])
                    if not lt:
                        idx += 1
                        continue
                    n2 = min(len(ptxt), len(lt), 14)
                    if n2 >= 2 and ptxt[:n2] == lt[:n2]:
                        matched = idx
                        break
                    idx += 1
                if matched is None:
                    anchors.append(None)
                    pitches.append(None)
                    continue
                anchors.append({
                    "y": round(lines_in[matched]["y0"] - ry0, 1),
                    "x": round(lines_in[matched]["x0"] - rx0, 1),
                })
                para_match.append((pi, matched))
                ptr = matched + 1

            for k, (pi, li) in enumerate(para_match):
                nxt_li = para_match[k + 1][1] if k + 1 < len(para_match) else total_lines
                seg = lines_in[li:nxt_li]
                if len(seg) >= 2:
                    ys = [r["y0"] for r in seg]
                    gaps = [ys[j + 1] - ys[j] for j in range(len(ys) - 1)
                            if 6 < ys[j + 1] - ys[j] < 60]
                    pitches.append(round(sorted(gaps)[len(gaps) // 2], 1) if gaps else None)
                else:
                    pitches.append(None)

            if any(a is not None for a in anchors):
                b["para_anchors"] = anchors
                b["para_pitches"] = pitches
    finally:
        try:
            doc.close()
        except Exception:  # noqa: BLE001
            pass


def build_richtext_template(docx_path, bg_out_path) -> Optional[Dict]:
    """一次构建新编辑模式所需的全部模板数据：
    PDF 校准几何 → 字体/段落格式 → 板块语义 → 照片框 → 无字装饰底图。
    任一环节失败返回 None（前端回退旧编辑模式）。"""
    docx_path = Path(docx_path)
    doc = Document(str(docx_path))
    sec = doc.sections[0]
    page_w = round(sec.page_width / EMU_PER_PX)
    page_h = round(sec.page_height / EMU_PER_PX)
    boxes = _collect_boxes(doc.element.body)

    # PDF 实际渲染位置校准（矩形包含式配对；未命中的退回 XML 几何 + 页边距）
    tmp_dir = Path(tempfile.mkdtemp(prefix="rtbuild_"))
    pdf_path = tmp_dir / "tpl.pdf"
    hit = set()
    ml = round(sec.left_margin / EMU_PER_PX)
    mt = round(sec.top_margin / EMU_PER_PX)
    try:
        if _to_pdf(docx_path, pdf_path):
            try:
                hit = _calibrate_containment(pdf_path, boxes, ml, mt)
            except Exception:  # noqa: BLE001
                hit = set()
    finally:
        for b in boxes:
            if b["id"] in hit:
                continue
            b["x"] += ml
            b["y"] += mt

    # 几何可信度检查：超过 30% 的可编辑盒子 PDF 校准失败（多为艺术字/
    # 组合形状/内联文本框等复杂结构），位置不可信，交给经典版渲染
    editable_boxes = [b for b in boxes if (b.get("text") or "").strip()]
    if editable_boxes:
        uncalibrated = sum(1 for b in editable_boxes if b["id"] not in hit)
        if uncalibrated > len(editable_boxes) * 0.3:
            print(f"[richtext] 校准失败率过高 "
                  f"({uncalibrated}/{len(editable_boxes)})，判定为复杂模板")
            shutil.rmtree(tmp_dir, ignore_errors=True)
            return None

    te._fix_degenerate_boxes(boxes, page_w, page_h)

    # 可编辑判定：有文字的盒子；无文字但够大的是待填内容区；
    # 纯隐藏文字（vanish）的盒子不进编辑层（留在底图/被遮盖，避免幽灵编辑区）。
    for b in boxes:
        if b.get("hidden"):
            b["editable"] = False
        else:
            b["editable"] = bool((b.get("text") or "").strip()) or (b["w"] >= 40 and b["h"] >= 14)

    # 隐藏冗余层排除：PDF 校准未命中的盒子里——
    # ① 每行文字都与某个其他盒子重复（被照片/色块遮盖的隐形副本）；
    # ② x 贴着页面左缘（x≤4，正常设计不会把正文放进出血区）。
    # 两类都不进编辑层，避免无法对位的幽灵编辑区。
    calibrated_ids = {b["id"] for b in boxes if b["id"] in hit}

    def _norm_lines(t):
        return [re.sub(r"\s+", "", ln) for ln in (t or "").splitlines() if ln.strip()]

    for b in boxes:
        if not b["editable"] or b["id"] in calibrated_ids:
            continue
        lines = _norm_lines(b["text"])
        pool = set()
        for other in boxes:
            if other["id"] != b["id"]:
                pool.update(_norm_lines(other["text"]))
        dup = bool(lines) and all(ln in pool for ln in lines)
        if dup or (b["x"] <= 4 and not lines == []):
            b["editable"] = False

    _annotate_roles(boxes)

    # 段落锚点：基于最终盒子几何，从同一份校准 PDF 提取每段实测位置
    # （必须在 tmp_dir 清理前执行）
    if pdf_path.exists():
        try:
            _extract_para_anchors(pdf_path, boxes)
        except Exception as exc:  # noqa: BLE001
            print(f"[WARN] 段落锚点提取失败: {exc}")

    photo_boxes = _detect_photos(doc.element.body)

    ok = make_blank_background(docx_path, bg_out_path, boxes)

    slim = []
    for b in boxes:
        if not b["editable"]:
            continue
        slim.append({
            "id": b["id"], "kind": b["kind"], "src": b["src"],
            "role": b["role"],
            "x": b["x"], "y": b["y"], "w": b["w"], "h": b["h"],
            "font": b["font"], "props": b["props"],
            "text": b["text"],
            "paras": [{"text": p["text"], "align": p["align"], "line": p["line"],
                       "font": _first_font([{"runs": p.get("runs") or []}])}
                      for p in b["paragraphs"]],
            "anchors": b.get("para_anchors"),
            "pitches": b.get("para_pitches"),
            "pitch": b.get("pitch"),
            "placeholders": b["placeholders"],
        })

    shutil.rmtree(tmp_dir, ignore_errors=True)

    return {"page_w": page_w, "page_h": page_h,
            "boxes": slim, "photo_boxes": photo_boxes,
            "bg_ok": bool(ok)}


def make_blank_background(docx_path, out_png, boxes: List[Dict]) -> bool:
    """生成「无字装饰底图」：副本里清空所有带文字的盒子（编辑层盒子 +
    被排除的隐藏冗余副本——否则副本仍会渲染出来形成文字双影），
    用 Word 渲染整页 PNG。装饰元素（形状/图标/照片）全部保留。"""
    tmp_dir = Path(tempfile.mkdtemp(prefix="rtblank_"))
    try:
        work = tmp_dir / "blank.docx"
        shutil.copyfile(docx_path, work)
        doc = Document(str(work))
        targets = _collect_edit_targets(doc.element.body)
        text_ids = {b["id"] for b in boxes if (b.get("text") or "").strip()}
        for idx, (_kind, node, _geo) in enumerate(targets):
            if idx in text_ids:
                te._write_plain_text(node, "")
        buf = io.BytesIO()
        doc.save(buf)
        work.write_bytes(buf.getvalue())

        import fitz
        pdf_path = tmp_dir / "blank.pdf"
        if not _to_pdf(work, pdf_path):
            return False
        pdf = fitz.open(str(pdf_path))
        try:
            pix = pdf[0].get_pixmap(matrix=fitz.Matrix(2.0, 2.0))
            Path(out_png).parent.mkdir(parents=True, exist_ok=True)
            pix.save(str(out_png))
        finally:
            pdf.close()
        return True
    except Exception:  # noqa: BLE001
        return False
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


# ═══════════════════ 富文本写回（导出用） ═══════════════════

_COLOR_RE = re.compile(r"^#?[0-9A-Fa-f]{6}$")


def _rpr_xml(run: Dict, base_font: Dict) -> Optional[OxmlElement]:
    rpr = OxmlElement("w:rPr")
    fonts = run.get("font") or base_font.get("family")
    if fonts:
        rf = OxmlElement("w:rFonts")
        rf.set(qn("w:eastAsia"), fonts)
        rf.set(qn("w:ascii"), fonts)
        rf.set(qn("w:hAnsi"), fonts)
        rpr.append(rf)
    if run.get("bold"):
        rpr.append(OxmlElement("w:b"))
    if run.get("italic"):
        rpr.append(OxmlElement("w:i"))
    if run.get("underline"):
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
    size = run.get("size") or base_font.get("size")
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(round(float(size) * 2))))
        rpr.append(sz)
    color = run.get("color") or base_font.get("color")
    if color and _COLOR_RE.match(color):
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color.lstrip("#").upper())
        rpr.append(c)
    return rpr if len(rpr) else None


def _ppr_xml(align: Optional[str], line: Optional[float], base_ppr) -> Optional[OxmlElement]:
    if base_ppr is not None:
        # 保留模板原段落格式；仅当用户显式改动时覆盖对齐/行距
        ppr = te.deepcopy(base_ppr)
        if align:
            jc = ppr.find(qn("w:jc"))
            if jc is None:
                jc = OxmlElement("w:jc")
                ppr.append(jc)
            jc.set(qn("w:val"), align)
        if line:
            sp = ppr.find(qn("w:spacing"))
            if sp is None:
                sp = OxmlElement("w:spacing")
                ppr.append(sp)
            sp.set(qn("w:line"), str(int(round(line * 240))))
            sp.set(qn("w:lineRule"), "auto")
        return ppr
    ppr = OxmlElement("w:pPr")
    if align and align != "left":
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), align)
        ppr.append(jc)
    if line:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:line"), str(int(round(line * 240))))
        sp.set(qn("w:lineRule"), "auto")
        ppr.append(sp)
    return ppr if len(ppr) else None


def apply_rich_texts(docx_path, edits: Dict) -> bytes:
    """把编辑器序列化的富文本写回模板（保留模板格式，支持段/字级覆盖）。

    edits: {box_id: {"paragraphs": [{"align": "left|center|right",
                                     "line": 1.5|null,
                                     "runs": [{"text", "bold", "italic",
                                               "underline", "size", "color"}]}]}}
    内容变长的盒子自动开启「缩小字体填充」防止溢出跑版。"""
    doc = Document(str(docx_path))
    body = doc.element.body
    targets = _collect_edit_targets(body)
    for box_id, payload in (edits or {}).items():
        try:
            idx = int(box_id)
        except (TypeError, ValueError):
            continue
        if idx < 0 or idx >= len(targets):
            continue
        kind, node, _geo = targets[idx]
        paras = payload.get("paragraphs") if isinstance(payload, dict) else None
        if not paras:
            continue
        # 基准格式：取写回前的第一段首 run（模板原格式）
        orig_paras = node.findall(qn("w:p"))
        base_font = {"family": "微软雅黑", "size": 12, "color": "#333333"}
        if orig_paras:
            info = _paragraph_info(orig_paras[0])
            for r in info["runs"]:
                if (r.get("text") or "").strip():
                    base_font = {"family": r.get("font") or "微软雅黑",
                                 "size": r.get("size") or 12,
                                 "color": r.get("color") or "#333333"}
                    break
        base_ppr = orig_paras[0].find(qn("w:pPr")) if orig_paras else None

        new_text_len = sum(len(r.get("text") or "")
                           for p in paras for r in (p.get("runs") or []))
        if kind == "textbox" and new_text_len > len(_node_text(node)):
            te._ensure_shrink_on_overflow(node)

        for p in list(orig_paras):
            node.remove(p)
        for para in paras:
            p_elem = OxmlElement("w:p")
            ppr = _ppr_xml(para.get("align"), para.get("line"), base_ppr)
            if ppr is not None:
                p_elem.append(ppr)
            runs = para.get("runs") or []
            if not runs:
                node.append(p_elem)
                continue
            for run in runs:
                text = run.get("text")
                if not text:
                    continue
                r = OxmlElement("w:r")
                rpr = _rpr_xml(run, base_font)
                if rpr is not None:
                    r.append(rpr)
                t = OxmlElement("w:t")
                t.set(qn("xml:space"), "preserve")
                t.text = text
                r.append(t)
                p_elem.append(r)
            node.append(p_elem)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ═══════════════════ 照片替换 ═══════════════════

def replace_photo_bytes(docx_bytes: bytes, rid: str, photo_bytes: bytes) -> bytes:
    """把模板里 rid 对应的图片部件替换为 photo_bytes。
    新照片按原图纵横比居中裁剪，保证模板内的裁剪形状（圆形等）不变形。"""
    from PIL import Image

    doc = Document(io.BytesIO(docx_bytes))
    part = doc.part.related_parts.get(rid)
    if part is None:
        raise KeyError(f"图片关系 {rid} 不存在")

    # 原图纵横比
    try:
        orig = part.image
        orig_w, orig_h = orig.px_width, orig.px_height
    except Exception:  # noqa: BLE001
        orig_w, orig_h = 3, 4

    img = Image.open(io.BytesIO(photo_bytes)).convert("RGB")
    target_ratio = orig_w / max(1, orig_h)
    w, h = img.size
    cur_ratio = w / max(1, h)
    if cur_ratio > target_ratio:      # 过宽 → 裁两侧
        new_w = int(h * target_ratio)
        left = (w - new_w) // 2
        img = img.crop((left, 0, left + new_w, h))
    else:                             # 过高 → 裁上下
        new_h = int(w / target_ratio)
        top = (h - new_h) // 2
        img = img.crop((0, top, w, top + new_h))
    if img.width > 800:
        img = img.resize((800, int(800 * img.height / img.width)), Image.LANCZOS)

    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=90)
    part._blob = buf.getvalue()       # ImagePart 保存时取 _blob
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
