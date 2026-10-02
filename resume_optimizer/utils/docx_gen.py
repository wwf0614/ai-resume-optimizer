"""从文本构建简单格式的 Word 简历文档。"""
import re
from pathlib import Path
from typing import Union

from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_LINE_SPACING


def create_resume_docx(text: str, output_path: Union[str, Path]) -> Path:
    """将纯文本简历按行解析为 Word 文档。

    - 被【】包裹的行 → 标题（加粗 14pt，段前间距 12pt）
    - 其他行 → 普通段落（11pt，1.5 倍行距）
    - 页面边距：上下 2cm，左右 2.5cm
    """
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue

        m = re.match(r"^【(.+?)】$", stripped)
        if m:
            p = doc.add_paragraph()
            run = p.add_run(f"【{m.group(1)}】")
            run.bold = True
            run.font.size = Pt(14)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
        else:
            p = doc.add_paragraph(stripped)
            p.style = doc.styles["Normal"]
            run = p.runs[0] if p.runs else None
            if run:
                run.font.size = Pt(11)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            p.paragraph_format.first_line_indent = Pt(0)

    doc.save(str(output_path))
    return Path(output_path)