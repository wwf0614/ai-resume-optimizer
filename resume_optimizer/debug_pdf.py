"""模拟 pdf2docx 输出：内容全在表格中，测试修复后能否识别区块。"""
from docx import Document

doc = Document()
# 模拟 pdf2docx：所有内容放在一个表格里
table = doc.add_table(rows=0, cols=1)

def add_row(text, bold=False):
    cell = table.add_row().cells[0]
    para = cell.paragraphs[0]
    run = para.add_run(text)
    run.bold = bold
    return para

add_row("张三 | 电话：13800000000 | 邮箱：zhangsan@example.com")
add_row("教育背景", bold=True)
add_row("2013.09 - 2017.06 某大学 计算机科学与技术 本科")
add_row("主修课程：数据结构、操作系统、数据库原理")
add_row("实习经历", bold=True)
add_row("2016.07 - 2016.09 某互联网公司 后端开发实习生")
add_row("负责后端接口的开发与测试。")
add_row("协助团队完成数据库迁移工作。")
add_row("校园经历", bold=True)
add_row("2015.09 - 2016.06 计算机协会 技术部部长")
add_row("组织协会技术沙龙活动，参与人数约50人。")
add_row("协调部门成员完成活动策划与执行。")
add_row("专业技能", bold=True)
add_row("熟悉Java、Python编程语言")
add_row("了解MySQL数据库和Redis缓存")
add_row("证书", bold=True)
add_row("CET-6 大学英语六级")
add_row("自我评价", bold=True)
add_row("本人性格开朗，善于沟通，有较强的学习能力和团队合作精神。")

doc.save(r"D:\DeepseekV4\resume_optimizer\test_table.docx")
print("ok")

# 测试识别
from utils.editor import _collect_all_paragraphs, _find_sections
doc2 = Document(r"D:\DeepseekV4\resume_optimizer\test_table.docx")
all_paras = _collect_all_paragraphs(doc2)
print(f"doc.paragraphs: {len(doc2.paragraphs)}, all_paras: {len(all_paras)}, tables: {len(doc2.tables)}")
print("\n段落列表:")
for i, p in enumerate(all_paras):
    print(f"  [{i}] {repr(p.text.strip()[:60])}")

sections = _find_sections(all_paras)
print(f"\n识别到区块: {len(sections)}")
for s, e, title, stype in sections:
    print(f"  [{s}-{e}] type={stype} title={title}")
