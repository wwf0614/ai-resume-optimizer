# -*- coding: utf-8 -*-
"""探查模板占位符与板块结构。"""
import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
sys.path.insert(0, os.getcwd())
from utils import db, filler, section_mapper
from docx import Document
from docx.oxml.ns import qn

for tid in (2, 11, 8):
    tpl = db.get_template(tid)
    fp = os.path.join(os.getcwd(), tpl["file_path"].replace("/", os.sep))
    doc = Document(fp)
    ph = set()
    for p in doc.element.iter(qn("w:p")):
        for m in filler.PLACEHOLDER_RE.finditer(p.text or ""):
            ph.add(m.group(1).strip())
    filler._remove_fallback_duplicates(doc)
    m = section_mapper.detect_mapping(doc)
    print("===== t%d %s =====" % (tid, tpl["name"]))
    print("占位符:", sorted(ph))
    print("板块:", [(s["key"], s["boxes"]) for s in m["sections"]])
    print("字段:", [f["key"] for f in m["fields"]])
    tx = doc.element.findall(".//" + qn("w:txbxContent"))
    print("表格数:", len(doc.tables))
    for i, tb in enumerate(doc.tables):
        print("  表格%d: %d行x%d列" % (i, len(tb.rows), len(tb.columns)))
        for row in tb.rows[:3]:
            print("    ", " | ".join(c.text.strip()[:12] for c in row.cells))
    print()
